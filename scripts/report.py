# -*- coding: utf-8 -*-
"""组价表报表生成器：清单+定额组价 → 规范 Excel（融合两种参考范例）。

输出三张表：
  1. 「组价表」——清单→定额→资源 三级层次（截图范例）：序号/清单编号/定额编号/名称/说明/
     单位/定额消耗/系数/定额工程量/含量/单价/合价/出处，尾部含「编制说明」块。
  2. 「综合单价分析表」——PDF 范例列式：数量/综合单价/人工费/材料费/机械费/管理费/利润/合价。
  3. 「编制说明与组价思路」——依据、范围、逐条透明计算过程、风险与价格取定。

用法：python report.py <组价任务.json> [-o 输出.xlsx]
      （须 openpyxl；运行环境 python -m pip install openpyxl）
"""
import argparse
import json
import os
import re
import sys
from datetime import date

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
from quota import BOOK_PROF, DEFAULT_FEE, QuotaDB              # noqa: E402

from openpyxl import Workbook                               # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

THIN = Side(style="thin", color="8EA9DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HEAD = PatternFill("solid", fgColor="4472C4")
FILL_SEC = PatternFill("solid", fgColor="C55A11")
FILL_BILL = PatternFill("solid", fgColor="D9E1F2")
FILL_QUOTA = PatternFill("solid", fgColor="E2EFDA")
FILL_RES = PatternFill("solid", fgColor="FFFFFF")
FILL_UNPRICED = PatternFill("solid", fgColor="FFF2CC")
FILL_TOTAL = PatternFill("solid", fgColor="FCE4D6")
F_TITLE = Font(name="微软雅黑", size=14, bold=True)
F_HEAD = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
F_BODY = Font(name="微软雅黑", size=9)
F_BOLD = Font(name="微软雅黑", size=9, bold=True)
F_NOTE = Font(name="微软雅黑", size=8, color="555555")

COLS = ["序号", "清单编号/定额编号", "项目名称/子目名称/资源名称", "项目特征/工作内容/说明",
        "单位", "定额消耗", "系数", "定额工程量", "含量", "单价(元)", "合价(元)", "出处/备注"]
WID = [4, 16, 30, 42, 6, 9, 6, 10, 9, 10, 11, 42]
AN_COLS = ["序号", "项目编码", "项目名称", "单位", "数量", "综合单价(元)",
           "人工费", "材料费", "机械费", "管理费", "利润", "合价(元)"]
AN_WID = [5, 15, 34, 6, 9, 12, 10, 10, 10, 10, 9, 12]


def fmt(v, n=2):
    return "" if v is None else round(v + 0.0, n)


def parse_rate(pr):
    """从 price 结果反推 管理费率/利润率（基数=人工+机械）。"""
    base = (pr["汇总"]["人工费"] or 0) + (pr["汇总"]["机械费"] or 0)
    if base <= 0:
        return 0.0, 0.0
    return (pr["汇总"]["管理费"] or 0) / base, (pr["汇总"]["利润"] or 0) / base


def compute(q, spec):
    """把 spec 算成带层次的行树 + 全部透明计算参数。"""
    head = spec.get("工程") or {}
    region, period = head.get("地区", "雄安新区"), head.get("期次", "2026-08")
    tax_mode, fee_cat = head.get("税制", "一般计税"), head.get("类别")
    tree, analyses = [], []
    used_books = set()
    for li, ln in enumerate(spec.get("清单") or [], 1):
        qty = float(ln.get("工程量") or 0)
        lnode = {"kind": "bill", "no": ln.get("编号", str(li)), "code": ln.get("清单编码", ""),
                 "name": ln.get("名称", ""), "desc": ln.get("特征", ""), "unit": ln.get("单位", ""),
                 "qty": qty, "note": ln.get("出处", ""), "subs": [], "chain": []}
        sum_sub = 0.0
        comp = {"人工费": 0.0, "材料费": 0.0, "机械费": 0.0, "管理费": 0.0, "利润": 0.0}
        for s in ln.get("子目") or []:
            if s.get("未计价主材") and not s.get("编号"):
                # 暂估/未计价主材行（无定额编号）：只列资源行，合价可带暂估价
                p0 = s.get("单价")
                n0 = s.get("数量", qty * float(s.get("含量", 1.0)))
                ssum = (p0 * n0) if (p0 is not None and n0 is not None) else None
                sub0 = {"kind": "quota", "code": "（未计价主材/暂估）",
                        "name": s.get("名称") or s["未计价主材"], "unit": s.get("单位", ln.get("单位", "")),
                        "coef": None, "Q": n0, "qty_unit": p0, "sum": ssum or 0.0,
                        "c": s.get("含量"), "book": "",
                        "note": s.get("说明", "暂估价/未计价主材，按招标文件计入"),
                        "src": "★暂估/未计价，结算按实调整",
                        "人工费": 0, "材料费": ssum or 0, "机械费": 0, "管理费": 0, "利润": 0,
                        "res": [{"kind": "res", "name": "√" + s["未计价主材"] + "（★暂估/未计价）",
                                 "unit": s.get("单位", ""), "u": None, "coef": None, "Q": n0,
                                 "total": n0, "p": p0, "sum": ssum,
                                 "src": "★暂估/未计价，按招标文件计入", "unpriced": True}]}
                lnode["subs"].append(sub0)
                sum_sub += sub0["sum"]
                comp["材料费"] += ssum or 0
                lnode["chain"].append(
                    f"· 子目（未计价主材/暂估）{sub0['name']}：数量 {n0:g}"
                    + (f"，暂估单价 {p0:.2f}，合价 {ssum:.2f} 元" if ssum is not None else "，单价待定"))
                continue
            book, code = s.get("册"), str(s.get("编号"))
            used_books.add(book)
            prof = s.get("专业") or BOOK_PROF.get(book, book)          # 册→取费专业（安装补充→安装）
            cat = s.get("类别") or DEFAULT_FEE.get(prof) or fee_cat    # 各册默认取费类别
            pr = q.price(code, book, region, period, prof, cat,
                         s.get("税制") or tax_mode)
            if "错误" in pr:
                lnode["subs"].append({"kind": "quota", "name": f"{book} {code}｜{pr['错误']}",
                                      "note": "定位失败", "err": True})
                continue
            c = float(s.get("含量", 1.0))
            f = float(s.get("人工系数", s.get("系数", 1.0)) or 1.0)
            Q = qty * c                                  # 定额工程量 = 清单量 × 含量
            r_mg, r_lr = parse_rate(pr)
            lab0 = pr["汇总"]["人工费"] or 0.0
            mat = pr["汇总"]["材料费"] or 0.0
            mac = pr["汇总"]["机械费"] or 0.0
            lab = lab0 * f
            gl = (lab + mac) * r_mg
            lr = (lab + mac) * r_lr
            unit_price = lab + mat + mac + gl + lr       # 定额单价（含系数）
            sub = {"kind": "quota",
                   "code": f"{book} {code}",
                   "name": pr["子目"]["名称"], "unit": pr["子目"]["单位"],
                   "coef": f if f != 1 else None, "Q": Q, "qty_unit": unit_price,
                   "sum": unit_price * Q, "c": c, "book": book,
                   "人工费": lab, "材料费": mat, "机械费": mac,
                   "管理费": gl, "利润": lr,
                   "note": (s.get("说明") or ""),
                   "src": (f"{book}册 {code}｜{(pr['子目'] or {}).get('标准号') or ''}"
                           f"｜原书P{pr['子目'].get('页码', '?')}"),
                   "res": []}
            it = q.con.execute(
                "SELECT b.standard_no,i.page_no FROM item i JOIN book b ON b.book_id=i.book_id"
                " WHERE i.book_id=? AND i.code=?", (book, code)).fetchone()
            if it:
                sub["src"] = f"{book}册 {code}｜{it['standard_no']}｜原书P{it['page_no']}"
            for d in pr["明细"]:
                u = d["消耗量"] or 0.0
                coef_r = f if (d["类别"] == "人工" and f != 1) else 1.0
                total = u * coef_r * Q
                amt = d["金额"] or 0.0
                # %类行（其他材料1.5%）：单价列是「基数」、金额=基数×%——直接用金额推有效单价，
                # 否则合价会把基数当单价乘（4143→41.43 的错）
                p_disp = (amt / u) if u else (d["单价"] or 0.0)
                sub["res"].append({"kind": "res", "name": ("√" + d["名称"]),
                                   "unit": d["单位"], "u": u,
                                   "coef": coef_r if coef_r != 1 else None,
                                   "Q": Q, "total": total, "p": p_disp,
                                   "sum": amt * coef_r * Q,
                                   "src": d.get("价格出处") or "",
                                   "unpriced": "未计价" in (d.get("价格出处") or "")
                                   or (d["单价"] in (None, 0))})
            # 管理费/利润以资源行列出：(人工+机械)×费率，让「资源行合计=定额行合价」严格闭合
            if gl or lr:
                sub["res"].append({
                    "kind": "res",
                    "name": f"√管理费（(人工+机械)×{r_mg * 100:.1f}%）",
                    "unit": "元", "u": None, "coef": None, "Q": Q, "total": None,
                    "p": gl, "sum": gl * Q, "src": f"费率取自{prof}·{cat}", "unpriced": False})
                sub["res"].append({
                    "kind": "res",
                    "name": f"√利润（(人工+机械)×{r_lr * 100:.1f}%）",
                    "unit": "元", "u": None, "coef": None, "Q": Q, "total": None,
                    "p": lr, "sum": lr * Q, "src": f"费率取自{prof}·{cat}", "unpriced": False})
            if s.get("未计价主材"):
                sub["res"].append({"kind": "res", "name": "√" + s["未计价主材"],
                                   "unit": s.get("单位", ""), "u": None, "Q": Q,
                                   "total": s.get("数量"), "p": None, "sum": None,
                                   "src": "★按主材/暂估，按招标文件计入", "unpriced": True})
            lnode["subs"].append(sub)
            sum_sub += sub["sum"]
            for k, v in (("人工费", lab * Q), ("材料费", mat * Q), ("机械费", mac * Q),
                         ("管理费", gl * Q), ("利润", lr * Q)):
                comp[k] += v
            lnode["chain"].append(
                f"· 子目 {sub['code']} {sub['name']}：含量 {c:g} → 定额工程量 = {qty:g}×{c:g} = {Q:.3f}"
                + (f"（人工系数 {f:g}）" if f != 1 else "")
                + f"\n    资源消耗 = 单位消耗 × {(f if f != 1 else 1):g} × {Q:.3f}；"
                + f"定额单价 = 人工{lab:.2f} + 材料{mat:.2f} + 机械{mac:.2f}"
                + f" + 管理费{gl:.2f} + 利润{lr:.2f} = {unit_price:.2f} 元/{sub['unit']}"
                + f"\n    定额合价 = {unit_price:.2f} × {Q:.3f} = {sub['sum']:.2f} 元")
        lnode["unit_price"] = (sum_sub / qty) if qty else 0.0
        lnode["sum"] = sum_sub
        lnode["comp"] = {k: (v / qty if qty else 0) for k, v in comp.items()}
        lnode["comp_sum"] = comp
        lnode["chain"].insert(0, f"「{lnode['code']} {lnode['name']}」清单工程量 {qty:g} {lnode['unit']}")
        lnode["chain"].append(
            f"· 清单综合单价 = Σ(定额单价×含量) = {lnode['unit_price']:.2f} 元/{lnode['unit']}；"
            f"合价 = {lnode['unit_price']:.2f} × {qty:g} = {sum_sub:.2f} 元")
        analyses.append(lnode)
        tree.append(lnode)
    return tree, analyses, used_books


# ---------------- 渲染 ----------------
def put(ws, r, c, v, font=F_BODY, fill=None, align=None, num=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align or Alignment(vertical="center", wrap_text=True)
    if border:
        cell.border = BORDER
    if num:
        cell.number_format = num
    return cell


def render_group(ws, r0, spec, tree, q):
    head = spec.get("工程") or {}
    put(ws, r0, 1, f"{head.get('名称', '工程')} 工程量清单与定额组价表", F_TITLE,
        align=Alignment(horizontal="center", vertical="center"), border=False)
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=12)
    ws.row_dimensions[r0].height = 26
    r = r0 + 1
    info = spec.get("编制信息") or {}
    lines = info.get("表头依据") or [
        "清单：GB/T 50500-2024《建设工程工程量清单计价标准》+ GB/T 50854-2024"
        "《房屋建筑与装饰工程工程量计算标准》+ 河北省《建设工程工程量清单计价指南》(冀建建市函〔2025〕154号)",
        "定额：河北省《房屋建筑与装饰工程消耗量标准》DB13(J)/T 8510/8511-2022、"
        "《房屋修缮工程消耗量标准》DB13(J)/T 8557-2022 等（含 14 号勘误、16 号降水调整）",
        f"价格：依据《河北建设工程造价信息》({head.get('期次', '2026-08')}期{head.get('地区', '雄安新区')})"
        f"及河北省 2023 基期价；{info.get('编制人', '')}"
        f"（制表：{date.today().isoformat()}）",
    ]
    for t in lines:
        put(ws, r, 1, t, F_NOTE, border=False)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        r += 1
    for j, h in enumerate(COLS, 1):
        put(ws, r, j, h, F_HEAD, FILL_HEAD,
            Alignment(horizontal="center", vertical="center", wrap_text=True))
    ws.row_dimensions[r].height = 30
    r += 1
    sec = info.get("分部标题", "一、分部分项工程项目清单")
    put(ws, r, 1, sec, F_BOLD, FILL_SEC)
    for j in range(2, 13):
        put(ws, r, j, None, F_BOLD, FILL_SEC)
    r += 1
    total = 0.0
    for node in tree:
        total += node["sum"]
        put(ws, r, 1, node["no"], F_BOLD, FILL_BILL)
        put(ws, r, 2, node["code"], F_BOLD, FILL_BILL)
        put(ws, r, 3, node["name"], F_BOLD, FILL_BILL)
        put(ws, r, 4, node["desc"], F_BODY, FILL_BILL)
        put(ws, r, 5, node["unit"], F_BODY, FILL_BILL, Alignment(horizontal="center"))
        put(ws, r, 6, fmt(node["qty"], 3), F_BOLD, FILL_BILL, num="0.000")
        for j in (7, 8, 9):
            put(ws, r, j, None, F_BODY, FILL_BILL)
        put(ws, r, 10, fmt(node["unit_price"]), F_BOLD, FILL_BILL, num="0.00")
        put(ws, r, 11, fmt(node["sum"]), F_BOLD, FILL_BILL, num="0.00")
        put(ws, r, 12, node["note"], F_NOTE, FILL_BILL)
        r += 1
        for sub in node["subs"]:
            if sub.get("err"):
                put(ws, r, 2, sub["name"], F_BODY, FILL_QUOTA)
                r += 1
                continue
            put(ws, r, 2, sub["code"], F_BODY, FILL_QUOTA)
            put(ws, r, 3, "「定额」" + sub["name"], F_BODY, FILL_QUOTA)
            put(ws, r, 4, sub["note"], F_BODY, FILL_QUOTA)
            put(ws, r, 5, sub["unit"], F_BODY, FILL_QUOTA, Alignment(horizontal="center"))
            put(ws, r, 6, None, F_BODY, FILL_QUOTA)
            put(ws, r, 7, fmt(sub["coef"], 2), F_BODY, FILL_QUOTA, num="0.##")
            put(ws, r, 8, fmt(sub["Q"], 3), F_BODY, FILL_QUOTA, num="0.000")
            put(ws, r, 9, fmt(sub["c"], 4), F_BODY, FILL_QUOTA, num="0.0000")
            put(ws, r, 10, fmt(sub["qty_unit"]), F_BODY, FILL_QUOTA, num="0.00")
            put(ws, r, 11, fmt(sub["sum"]), F_BODY, FILL_QUOTA, num="0.00")
            put(ws, r, 12, sub["src"], F_NOTE, FILL_QUOTA)
            r += 1
            for rs in sub["res"]:
                fill = FILL_UNPRICED if rs.get("unpriced") else FILL_RES
                put(ws, r, 3, rs["name"], F_BODY, fill)
                put(ws, r, 5, rs["unit"], F_BODY, fill, Alignment(horizontal="center"))
                put(ws, r, 6, fmt(rs["u"], 4), F_BODY, fill, num="0.0000")
                put(ws, r, 7, fmt(rs["coef"], 2), F_BODY, fill, num="0.##")
                put(ws, r, 8, fmt(rs["Q"], 3), F_BODY, fill, num="0.000")
                put(ws, r, 9, fmt(rs["total"], 4), F_BODY, fill, num="0.0000")
                put(ws, r, 10, fmt(rs["p"]), F_BODY, fill, num="0.00")
                put(ws, r, 11, fmt(rs["sum"]), F_BODY, fill, num="0.00")
                put(ws, r, 12, rs["src"], F_NOTE, fill)
                for j in (1, 2, 4):
                    put(ws, r, j, None, F_BODY, fill)
                r += 1
    put(ws, r, 1, "分部分项 + 措施项目 合计", F_BOLD, FILL_TOTAL)
    for j in range(2, 11):
        put(ws, r, j, None, F_BOLD, FILL_TOTAL)
    put(ws, r, 11, fmt(total), F_BOLD, FILL_TOTAL, num="0.00")
    put(ws, r, 12, None, F_BOLD, FILL_TOTAL)
    return r + 2


def render_analysis(ws, r0, tree):
    put(ws, r0, 1, "表1-19  分部分项工程量清单综合单价分析表", F_TITLE,
        align=Alignment(horizontal="center"), border=False)
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=12)
    r = r0 + 1
    put(ws, r, 1, "工程名称：" + "（见组价表）", F_NOTE, border=False)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    r += 1
    for j, h in enumerate(AN_COLS, 1):
        put(ws, r, j, h, F_HEAD, FILL_HEAD, Alignment(horizontal="center", wrap_text=True))
    r += 1
    for i, node in enumerate(tree, 1):
        put(ws, r, 1, i, F_BOLD, FILL_BILL)
        put(ws, r, 2, node["code"], F_BOLD, FILL_BILL)
        put(ws, r, 3, node["name"], F_BOLD, FILL_BILL)
        put(ws, r, 4, node["unit"], F_BODY, FILL_BILL, Alignment(horizontal="center"))
        put(ws, r, 5, fmt(node["qty"], 3), F_BOLD, FILL_BILL, num="0.000")
        put(ws, r, 6, fmt(node["unit_price"]), F_BOLD, FILL_BILL, num="0.00")
        for j, k in zip(range(7, 12), ("人工费", "材料费", "机械费", "管理费", "利润")):
            put(ws, r, j, fmt(node["comp"][k]), F_BODY, FILL_BILL, num="0.00")
        put(ws, r, 12, fmt(node["sum"]), F_BOLD, FILL_BILL, num="0.00")
        r += 1
        for j2, sub in enumerate([s for s in node["subs"] if not s.get("err")], 1):
            put(ws, r, 1, f"{i}.{j2}", F_BODY, FILL_QUOTA)
            put(ws, r, 2, sub["code"], F_BODY, FILL_QUOTA)
            put(ws, r, 3, sub["name"], F_BODY, FILL_QUOTA)
            put(ws, r, 4, sub["unit"], F_BODY, FILL_QUOTA, Alignment(horizontal="center"))
            put(ws, r, 5, fmt(sub["Q"], 3), F_BODY, FILL_QUOTA, num="0.000")
            put(ws, r, 6, fmt(sub["qty_unit"]), F_BODY, FILL_QUOTA, num="0.00")
            # 定额行组成按定额单位
            for j, k in zip(range(7, 12), ("人工费", "材料费", "机械费", "管理费", "利润")):
                v = sub.get(k)
                put(ws, r, j, fmt(v), F_BODY, FILL_QUOTA, num="0.00")
            put(ws, r, 12, fmt(sub["sum"]), F_BODY, FILL_QUOTA, num="0.00")
            r += 1
    return r


def render_notes(ws, r0, spec, analyses, used_books, q):
    put(ws, r0, 1, "编制说明", F_TITLE, align=Alignment(horizontal="left"), border=False)
    r = r0 + 1
    head = spec.get("工程") or {}
    stds = {}
    for b in sorted(used_books):
        row = q.con.execute("SELECT standard_no FROM book WHERE book_id=?", (b,)).fetchone()
        stds[b] = row[0] if row and row[0] else "（标准号待补）"
    std_line = "；".join(f"{b} {s}" for b, s in stds.items())
    items = [
        ("一、依据清单", (spec.get("编制信息") or {}).get("依据清单") or
         "GB/T 50500-2024《建设工程工程量清单计价标准》（2025-09-01 实施）；"
         "GB/T 50854-2024《房屋建筑与装饰工程工程量计算标准》；"
         "河北省《建设工程工程量清单计价指南》冀建建市函〔2025〕154 号（2026-01-01 施行）。"),
        ("二、依据定额", (spec.get("编制信息") or {}).get("依据定额") or
         f"河北省消耗量标准：{std_line}；含 14 号勘误（2025-02-13）、16 号降水工程子目调整（2026-08-01）。"),
        ("三、编制范围", (spec.get("编制信息") or {}).get("编制范围") or
         f"本表为{head.get('名称', '工程')}分部分项工程量清单与定额组价；"
         "措施项目、其他项目按合同约定与 05 号费用标准另行计取。"),
        ("四、组价思路（计算过程见第三表逐条展开）",
         "① 每条清单按项目特征与工作内容匹配定额子目，含量 = 1 个清单单位所需定额单位数"
         "（按清单单位与定额单位换算，如 清单 m / 定额 10m → 含量 0.1）；"
         "② 定额工程量 = 清单工程量 × 含量；"
         "③ 资源消耗量 = 定额单位消耗 × 换算系数（如保护性拆除人工×1.5）× 定额工程量；"
         "④ 资源合价 = 消耗量 × 资源单价（信息价/基期价，未计价主材单列并注明暂估）；"
         "⑤ 定额单价 = 人工 + 材料 + 机械 + 管理费 + 利润，其中管理费、利润以（人工+机械）为基数；"
         "⑥ 清单综合单价 = Σ(定额单价 × 含量)，合价 = 综合单价 × 清单工程量；"
         "⑦ 措施项目费 =（人工费+机械费）× 规定费率分专业计取；安全文明施工费按档位计取；"
         "增值税按税制计取（综合单价均为除税价）。"),
        ("五、施工措施项目", (spec.get("编制信息") or {}).get("措施") or
         "脚手架、垂直运输、冬雨季施工等按 05 号附件 1 措施项目费率自动计取；"
         "临时设施按施工组织设计/施工方案列入措施项目（清单第 7 项）。"),
        ("六、风险提示", (spec.get("编制信息") or {}).get("风险") or
         "单价风险范围按 GB/T 50500-2024 第 3.3 节执行，不得约定无限风险；"
         "材料价格涨跌超约定幅度时按合同调差条款执行；保护性拆除、超常规运距等已在说明中标注，"
         "漏项风险由清单缺陷规则（多列项/错漏项/特征不符/量偏差）划分责任。"),
        ("七、单价与价格取定", (spec.get("编制信息") or {}).get("价格") or
         f"人工单价按河北省 2023 基期（综合用工一/二/三类）；材料、机械按《河北建设工程造价信息》"
         f"{head.get('期次', '2026-08')}期{head.get('地区', '雄安新区')}价，缺项按 2023 基期价；"
         "未计价主材（带★）按招标文件暂估价列入，结算按实调整。"),
    ]
    for title, body in items:
        put(ws, r, 1, title, F_BOLD, border=False)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        r += 1
        put(ws, r, 1, body, F_BODY, border=False)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        ws.row_dimensions[r].height = max(15, 13 * (len(body) // 90 + 1))
        r += 2
    put(ws, r, 1, "逐条组价计算过程", F_BOLD, border=False)
    r += 1
    for node in analyses:
        for line in node["chain"]:
            put(ws, r, 1, line, F_BODY if line.startswith("·") else F_BOLD, border=False)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
            r += 1
        r += 1
    return r


def main():
    ap = argparse.ArgumentParser(description="组价表报表生成器")
    ap.add_argument("spec", help="组价任务 JSON（结构见 清单示例_组价表.json）")
    ap.add_argument("-o", "--out", help="输出 xlsx 路径（默认 同名_组价表.xlsx）")
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8"))
    q = QuotaDB()
    tree, analyses, used_books = compute(q, spec)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "组价表"
    for j, w in enumerate(WID, 1):
        ws1.column_dimensions[chr(64 + j) if j <= 26 else "A" + chr(64 + j - 26)].width = w
    r = render_group(ws1, 1, spec, tree, q)
    render_notes(ws1, r, spec, analyses, used_books, q)   # 编制说明放组价表尾部（截图范例）

    ws2 = wb.create_sheet("综合单价分析表")
    for j, w in enumerate(AN_WID, 1):
        ws2.column_dimensions[chr(64 + j)].width = w
    render_analysis(ws2, 1, tree)

    ws3 = wb.create_sheet("编制说明与组价思路")
    ws3.column_dimensions["A"].width = 120
    render_notes(ws3, 1, spec, analyses, used_books, q)

    out = a.out or (os.path.splitext(a.spec)[0] + "_组价表.xlsx")
    wb.save(out)
    total = sum(n["sum"] for n in tree)
    print(f"WROTE {out}｜清单 {len(tree)} 条｜分部分项合计 {total:,.2f} 元")


if __name__ == "__main__":
    main()
